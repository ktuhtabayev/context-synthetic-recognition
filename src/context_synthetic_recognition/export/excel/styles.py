"""The style palette of the Excel experiment, by meaning.

Every cell style of the author's workbook has a name here (``title``, ``bar``, ``header``,
``label``, ``class_``, ``synthetic``, ``latent_num``, ``warn1_num``, …), so the mirror uses the same
colour semantics: dark blue title bars, grey-blue column headers, grey object labels, blue
quantitative and yellow nominal values, red class labels, green synthetic features, the orange
⚠1 and pink ⚠2 columns of the two template deviations, and so on (see the workbook's *Overview*,
"Colour legend").

Colours are RGB hex without alpha; the font is Calibri, 12 pt unless stated.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook import Workbook

FONT_NAME = "Calibri"
DEFAULT_SIZE = 12

NUMBER = "0.0000"
"""Real numbers show four decimals (the stored value keeps every digit)."""
PERCENT = "0.0%"
K_HEADER = '"k = "0'
"""A permitted k as a column header: ``k = 3``."""
YES_NO = '"Yes";"Yes";"No"'
EXECUTED = '"executed";"executed";"not executed"'
RANK_OR_DASH = '[=999]"—";0'
"""The rank 999 of an excluded object is shown as "—"."""

_SIDES = {"t": Side(style="thin"), "m": Side(style="medium"), "-": Side()}


@dataclass(frozen=True)
class Style:
    """A cell style: fill, font, number format, alignment and border."""

    fill: str | None = None
    """Background colour (RGB hex)."""
    size: int = DEFAULT_SIZE
    bold: bool = False
    italic: bool = False
    color: str | None = None
    """Font colour (RGB hex); the default text colour if ``None``."""
    underline: bool = False
    fmt: str = "General"
    """Number format."""
    h: str | None = None
    """Horizontal alignment: ``left`` or ``center``."""
    v: bool = False
    """Centre vertically."""
    wrap: bool = False
    indent: bool = False
    """One step of indent (left-aligned text)."""
    border: str = "----"
    """Left, right, top, bottom: ``t`` thin, ``m`` medium, ``-`` none."""

    def apply(self, cell: Any) -> None:
        """Set this style on an openpyxl cell."""
        cell.font = Font(
            name=FONT_NAME,
            size=self.size,
            bold=self.bold,
            italic=self.italic,
            color=None if self.color is None else "FF" + self.color,
            underline="single" if self.underline else None,
        )
        if self.fill is not None:
            cell.fill = PatternFill("solid", fgColor="FF" + self.fill)
        cell.number_format = self.fmt
        cell.alignment = Alignment(
            horizontal=self.h,
            vertical="center" if self.v else None,
            wrap_text=self.wrap,
            indent=1 if self.indent else 0,
        )
        left, right, top, bottom = (_SIDES[code] for code in self.border)
        cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def _box(**kwargs: Any) -> Style:
    """A bordered, vertically centred cell — nearly every cell of the workbook."""
    return Style(v=True, border="tttt", **kwargs)


def _centred(**kwargs: Any) -> Style:
    return _box(h="center", **kwargs)


def _left(**kwargs: Any) -> Style:
    return _box(h="left", indent=True, **kwargs)


DARK_BLUE = "1F3864"
BLUE = "2F5597"
HEADER = "D6DCE4"
LABEL = "EDEDED"
KEY = "E4DFEC"
QUANTITATIVE = "DDEBF7"
NOMINAL = "FFF2CC"
CLASS = "F4CCCC"
CLASS_INK = "9C0006"
SYNTHETIC = "E2EFDA"
LATENT_HEADER = "FFE699"
INITIAL = "DFF3EA"
LATENT = "B7DEA8"
GOOD = "C6EFCE"
GOOD_INK = "006100"
BAD = "FFC7CE"
BAD_INK = "9C0006"
WARN1 = "FCE4D6"
WARN1_STRONG = "F8CBAD"
WARN1_INK = "833C0B"
WARN2 = "FFE1F0"
WARN2_STRONG = "FFC7E0"
WARN2_INK = "9C0055"
FLAG = "D9E1F2"
SHADE = "F2F2F2"
WHITE = "FFFFFF"

STYLES: dict[str, Style] = {
    # ------------------------------------------------------------ Overview
    "overview_title": Style(
        fill=DARK_BLUE, size=20, bold=True, color=WHITE, h="left", v=True, indent=True
    ),
    "overview_subtitle": Style(size=11, italic=True, color="595959", h="left", v=True, indent=True),
    "section": Style(size=13, bold=True, color=DARK_BLUE, h="left", v=True, border="---m"),
    "overview_head_center": _centred(fill=BLUE, size=11, bold=True, color=WHITE),
    "overview_head": _left(fill=BLUE, size=11, bold=True, color=WHITE),
    "overview_number": _centred(size=11),
    "link": _left(size=11, bold=True, color="0563C1", underline=True),
    "overview_text": _left(size=11, wrap=True),
    "overview_text_bold": _left(size=11, bold=True, wrap=True),
    "overview_spec": _left(fill="F7F9FC", size=10, color=DARK_BLUE, wrap=True),
    "overview_value": _centred(fill=SHADE, bold=True, color=DARK_BLUE),
    "overview_ref": _left(size=10, italic=True, color="595959"),
    "overview_percent": _centred(fill=SHADE, bold=True, color=DARK_BLUE, fmt=PERCENT),
    "overview_note": Style(size=10, color="404040", h="left", v=True, wrap=True, indent=True),
    "legend_title": _centred(fill=DARK_BLUE, size=11, bold=True, color=WHITE),
    "legend_header": _centred(fill=HEADER, size=11, bold=True),
    "legend_label": _centred(fill=LABEL, size=11, bold=True),
    "legend_quantitative": _centred(fill=QUANTITATIVE, size=11, bold=True),
    "legend_nominal": _centred(fill=NOMINAL, size=11, bold=True),
    "legend_class": _centred(fill=CLASS, size=11, bold=True, color=CLASS_INK),
    "legend_synthetic": _centred(fill=SYNTHETIC, size=11, bold=True),
    "legend_latent_header": _centred(fill=LATENT_HEADER, size=11, bold=True, color=DARK_BLUE),
    "legend_good": _centred(fill=GOOD, size=11, bold=True, color=GOOD_INK),
    "legend_initial": _centred(fill=INITIAL, size=11, bold=True),
    "legend_latent": _centred(fill=LATENT, size=11, bold=True),
    "legend_key": _centred(fill=KEY, size=11, bold=True, color="3F3151"),
    "legend_input": _centred(fill=WHITE, size=11, bold=True, color="0000CC"),
    "legend_warn1": _centred(fill=WARN1, size=11, bold=True, color=WARN1_INK),
    "legend_warn2": _centred(fill=WARN2, size=11, bold=True, color=WARN2_INK),
    # ------------------------------------------------------------ bars and headers
    "title": _centred(fill=DARK_BLUE, bold=True, color=WHITE),
    "bar": _centred(fill=BLUE, bold=True, color=WHITE),
    "bar_left": _left(fill=BLUE, bold=True, color=WHITE),
    "warn1_bar": _left(fill="C55A11", bold=True, color=WHITE),
    "warn2_bar": _left(fill="C00070", bold=True, color=WHITE),
    "tail": Style(size=11, border="--tt"),
    "tail_end": Style(size=11, border="-ttt"),
    "header": _centred(fill=HEADER, bold=True, italic=True, wrap=True),
    "header_plain": _centred(fill=HEADER, bold=True, italic=True),
    "header_open": Style(fill=HEADER, bold=True, italic=True, h="center", v=True, border="tt-t"),
    "header_left": _left(fill=HEADER, bold=True, italic=True),
    "header_k": _centred(fill=HEADER, bold=True, italic=True, fmt=K_HEADER),
    "executed": _centred(fill=HEADER, size=11, bold=True, italic=True, fmt=EXECUTED),
    "latent_header": _centred(fill=LATENT_HEADER, bold=True, italic=True, color=DARK_BLUE),
    # ------------------------------------------------------------ labels and captions
    "label": _centred(fill=LABEL, bold=True, italic=True),
    "caption": _left(fill=LABEL, bold=True, italic=True),
    "key": _left(fill=KEY, bold=True, italic=True, color="3F3151"),
    "key_center": _centred(fill=KEY, bold=True, italic=True, color="3F3151"),
    "value": _centred(fill=KEY, bold=True),
    "value_left": _left(fill=KEY, bold=True),
    "value_wrap": _left(fill=KEY, bold=True, wrap=True),
    "value_num": _centred(fill=KEY, bold=True, fmt=NUMBER),
    "value_yes_no": _centred(fill=KEY, bold=True, fmt=YES_NO),
    "status": _left(fill=SHADE, bold=True),
    "decision": _left(fill="F2DCEB", bold=True, color="7030A0"),
    "text": _left(fill=WHITE, size=11, color="404040"),
    "text_wrap": _left(fill=WHITE, size=11, color="404040", wrap=True),
    "note": Style(size=10, color="595959", h="left", v=True),
    # ------------------------------------------------------------ values
    "input": _centred(fill=WHITE, bold=True, color="0000CC"),
    "num": _centred(fill=WHITE, fmt=NUMBER),
    "integer": _centred(fill=WHITE),
    "small": _centred(fill=WHITE, size=11),
    "muted": _centred(fill=WHITE, color="595959"),
    "operator": _centred(fill=WHITE, bold=True, color="595959"),
    "yes_no": _centred(fill=WHITE, fmt=YES_NO),
    "rank": _centred(fill=WHITE, fmt=RANK_OR_DASH),
    "diagonal": _centred(fill=SHADE, color="7F7F7F"),
    "diagonal_num": _centred(fill=SHADE, color="7F7F7F", fmt=NUMBER),
    "statistic_label": _centred(fill=SHADE, bold=True, italic=True),
    "statistic": _centred(fill=SHADE, italic=True),
    "quantitative": _centred(fill=QUANTITATIVE),
    "quantitative_num": _centred(fill=QUANTITATIVE, fmt=NUMBER),
    "quantitative_flag": _centred(fill=QUANTITATIVE, bold=True, italic=True),
    "nominal": _centred(fill=NOMINAL),
    "nominal_flag": _centred(fill=NOMINAL, bold=True, italic=True),
    "flag": _centred(fill=FLAG, bold=True, italic=True),
    "size_note": Style(
        fill=FLAG, bold=True, italic=True, color=DARK_BLUE, h="center", border="tttt"
    ),
    "size_note_centered": _centred(fill=FLAG, bold=True, italic=True, color=DARK_BLUE),
    "class_": _centred(fill=CLASS, bold=True, italic=True, color=CLASS_INK),
    "synthetic": _centred(fill=SYNTHETIC, bold=True),
    "synthetic_num": _centred(fill=SYNTHETIC, bold=True, fmt=NUMBER),
    "initial_num": _centred(fill=INITIAL, fmt=NUMBER),
    "latent_num": _centred(fill=LATENT, fmt=NUMBER),
    "good": _centred(fill=GOOD, bold=True, color=GOOD_INK),
    "good_num": _centred(fill=GOOD, bold=True, color=GOOD_INK, fmt=NUMBER),
    # ------------------------------------------------------------ ⚠ the two template deviations
    "warn_note": _left(fill=WARN1, size=11, bold=True, color=WARN1_INK, wrap=True),
    "warn_text": _left(fill=WARN1, bold=True, color=WARN1_INK, wrap=True),
    "warn1_key": _left(fill=WARN1, bold=True, italic=True, color=WARN1_INK),
    "warn1_text": _left(fill=WARN1, size=11, bold=True, color=WARN1_INK),
    "warn1_header": _centred(
        fill=WARN1_STRONG, size=11, bold=True, italic=True, color=WARN1_INK, wrap=True
    ),
    "warn1_num": _centred(fill=WARN1, bold=True, color=WARN1_INK, fmt=NUMBER),
    "warn1_box": Style(
        fill=WARN1,
        size=11,
        bold=True,
        color=WARN1_INK,
        h="left",
        v=True,
        wrap=True,
        indent=True,
        border="mmmm",
    ),
    "warn2_key": _left(fill=WARN2, bold=True, italic=True, color=WARN2_INK),
    "warn2_text": _left(fill=WARN2, size=11, bold=True, color=WARN2_INK),
    "warn2_header": _centred(
        fill=WARN2_STRONG, size=11, bold=True, italic=True, color=WARN2_INK, wrap=True
    ),
    "warn2_num": _centred(fill=WARN2, bold=True, color=WARN2_INK, fmt=NUMBER),
    "warn2_latent": _centred(fill=WARN2_STRONG, bold=True, color=WARN2_INK, fmt=NUMBER),
    "warn2_box": Style(
        fill=WARN2,
        size=11,
        bold=True,
        color=WARN2_INK,
        h="left",
        v=True,
        wrap=True,
        indent=True,
        border="mmmm",
    ),
    # the cells of a merged note box keep only its outer (medium) border
    "box_top": Style(size=11, border="--m-"),
    "box_top_right": Style(size=11, border="-mm-"),
    "box_left": Style(size=11, border="m---"),
    "box_right": Style(size=11, border="-m--"),
    "box_bottom_left": Style(size=11, border="m--m"),
    "box_bottom": Style(size=11, border="---m"),
    "box_bottom_right": Style(size=11, border="-m-m"),
}
"""The palette: name → style."""

HIGHLIGHTS: dict[str, tuple[str, str]] = {
    "good": (GOOD, GOOD_INK),
    "bad": (BAD, BAD_INK),
    "neighbourhood": (LATENT_HEADER, DARK_BLUE),
}
"""Conditional highlights (fill, bold font colour): selected / same class / correct; wrong or
refused; the nested neighbourhoods of the permitted k."""

TAB_COLORS = {
    "info": "404040",
    "deviations": "833C0B",
    "data": "1F4E78",
    "quantitative": "2E75B6",
    "nominal": "BF8F00",
    "context": "7030A0",
    "synthetic": "C55A11",
    "hag": "548235",
    "meta": "2E75B6",
    "evaluation": "BF8F00",
    "checks": "7F7F7F",
}
"""Sheet-tab colours by stage."""


class StyleBook:
    """The palette prepared for one workbook: styles are resolved once and copied to cells.

    Setting font, fill, border and alignment on every cell costs four style look-ups per cell;
    copying a cell's resolved style record is an order of magnitude faster, which matters for the
    large sheets (neighbour blocks, HAG candidate blocks).
    """

    def __init__(self, workbook: Workbook) -> None:
        """Resolve every style of the palette in ``workbook``."""
        scratch = workbook.create_sheet("_styles")
        self._records: dict[str, Any] = {}
        for row, (name, style) in enumerate(STYLES.items(), start=1):
            cell = scratch.cell(row, 1)
            style.apply(cell)
            self._records[name] = copy(cell._style)
        workbook.remove(scratch)

    def set(self, cell: Any, name: str) -> None:
        """Give ``cell`` the style ``name``.

        Raises:
            KeyError: Unknown style name.
        """
        cell._style = copy(self._records[name])
