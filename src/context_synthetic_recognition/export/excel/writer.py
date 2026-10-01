"""Writing styled cells: the small vocabulary the sheet builders are written in.

A :class:`SheetWriter` wraps one worksheet of the mirror. Rows and columns are 1-based, as in
Excel (column 1 = A). Values are converted to what a cell can hold — numpy scalars to Python
numbers, NaN to the workbook's "—", booleans to 1/0 — and every write names a style of the palette
(:mod:`.styles`).
"""

from __future__ import annotations

import math
from collections.abc import Callable, Iterable, Sequence
from typing import Any

import numpy as np
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from context_synthetic_recognition.export.excel.styles import HIGHLIGHTS, TAB_COLORS, StyleBook
from context_synthetic_recognition.notation import DASH

WIDTH_PADDING = 0.77734375
"""Excel stores a column width of N characters as N + this padding (Calibri 11)."""
TITLE_HEIGHT = 20.0
DEFAULT_ROW_HEIGHT = 15.6
MAX_SHEET_NAME = 31
_FORBIDDEN_IN_SHEET_NAMES = "[]:*?/\\"

StyleChoice = str | Sequence[str] | Callable[[int, int], str]
"""One style for all cells, one per cell of a row/column, or a function of (row, column) offsets."""


class LayoutError(RuntimeError):
    """Two tables of a generated sheet overlap — a bug in a sheet builder."""


def cell_value(value: Any) -> Any:
    """A value as a cell holds it: booleans as 1/0, NaN as "—", ±∞ as text."""
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool | np.bool_):
        return int(value)
    if isinstance(value, int | np.integer):
        return int(value)
    number = float(value)
    if math.isnan(number):
        return DASH
    if math.isinf(number):
        return "+∞" if number > 0 else "−∞"
    return number


def column_letter(column: int) -> str:
    """1 → ``A``, 27 → ``AA``."""
    return str(get_column_letter(column))


def reference(row: int, column: int, rows: int = 1, columns: int = 1) -> str:
    """``B5`` or ``B5:G14`` for a block of ``rows`` × ``columns`` cells."""
    first = f"{column_letter(column)}{row}"
    if rows == 1 and columns == 1:
        return first
    return f"{first}:{column_letter(column + columns - 1)}{row + rows - 1}"


def sheet_name(text: str, taken: Iterable[str] = ()) -> str:
    r"""A valid, unique worksheet name: at most 31 characters, none of ``[]:*?/\``."""
    cleaned = "".join("·" if c in _FORBIDDEN_IN_SHEET_NAMES else c for c in text).strip("'")
    name = cleaned[:MAX_SHEET_NAME] or "Sheet"
    used = {t.lower() for t in taken}
    counter = 2
    while name.lower() in used:
        suffix = f" ({counter})"
        name = cleaned[: MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    return name


class SheetWriter:
    """One worksheet of the mirror."""

    def __init__(
        self,
        workbook: Workbook,
        styles: StyleBook,
        title: str,
        tab: str,
        *,
        freeze: str | None = "A3",
        check_overlaps: bool = False,
    ) -> None:
        """Create the sheet ``title`` with the workbook's page set-up and tab colour ``tab``.

        With ``check_overlaps`` a cell written twice raises :class:`LayoutError` — the tests use
        it to prove that the tables of a generated layout never run into each other.
        """
        self.ws: Worksheet = workbook.create_sheet(title)
        self.styles = styles
        self._written: set[tuple[int, int]] | None = set() if check_overlaps else None
        ws = self.ws
        ws.sheet_properties.tabColor = "FF" + TAB_COLORS[tab]
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        ws.sheet_format.defaultRowHeight = DEFAULT_ROW_HEIGHT
        ws.freeze_panes = freeze
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = 9  # A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = 0.7
        ws.page_margins.bottom = 0.6
        ws.page_margins.header = ws.page_margins.footer = 0.3
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.center.size = 9
        self.cells = 0
        """Cells written so far (for the size report)."""

    @property
    def title(self) -> str:
        """The sheet's name."""
        return str(self.ws.title)

    # ------------------------------------------------------------------ cells

    def _claim(self, row: int, column: int) -> Any:
        """The cell at (row, column), checked against earlier writes if overlaps are checked."""
        if self._written is not None:
            if (row, column) in self._written:
                raise LayoutError(
                    f"{self.title}!{column_letter(column)}{row} is written twice: two tables "
                    "of the layout overlap"
                )
            self._written.add((row, column))
        return self.ws.cell(row, column)

    def put(self, row: int, column: int, value: Any, style: str) -> None:
        """Write one cell."""
        cell = self._claim(row, column)
        cell.value = cell_value(value)
        self.styles.set(cell, style)
        self.cells += 1

    def row(self, row: int, column: int, values: Iterable[Any], style: StyleChoice) -> None:
        """Write values left to right from (``row``, ``column``)."""
        for offset, value in enumerate(values):
            self.put(row, column + offset, value, _pick(style, offset, 0, offset))

    def column(self, row: int, column: int, values: Iterable[Any], style: StyleChoice) -> None:
        """Write values top to bottom from (``row``, ``column``)."""
        for offset, value in enumerate(values):
            self.put(row + offset, column, value, _pick(style, offset, offset, 0))

    def block(self, row: int, column: int, matrix: Any, style: StyleChoice) -> None:
        """Write a matrix (rows × columns) with its top-left cell at (``row``, ``column``).

        ``style`` may be one name, one name per column, or a function of the (row, column) offsets.
        """
        for r, line in enumerate(matrix):
            for c, value in enumerate(line):
                self.put(row + r, column + c, value, _pick(style, c, r, c))

    def merge(
        self,
        row: int,
        column: int,
        width: int,
        value: Any,
        style: str,
        *,
        height: int = 1,
    ) -> None:
        """Write a value into a merged block; the covered cells keep the block's outer border."""
        self.put(row, column, value, style)
        if width == 1 and height == 1:
            return
        for r in range(row, row + height):
            for offset in range(0 if r > row else 1, width):
                cell = self._claim(r, column + offset)
                if height == 1:
                    self.styles.set(cell, "tail_end" if offset == width - 1 else "tail")
        self.join(row, column, width, height)

    def join(self, row: int, column: int, width: int, height: int = 1) -> None:
        """Merge a block of cells without touching their styles.

        ``Worksheet.merge_cells`` rebuilds the border of every covered cell, which is slow for the
        wide bars of a large sheet; the builders style the covered cells themselves.
        """
        self.ws.merged_cells.add(reference(row, column, height, width))

    def title_bar(self, text: str, width: int, *, height: float = TITLE_HEIGHT) -> None:
        """The sheet's title bar in row 1 across ``width`` columns."""
        self.merge(1, 1, width, text, "title")
        self.ws.row_dimensions[1].height = height

    def bar(self, row: int, column: int, width: int, text: str, style: str = "bar") -> None:
        """A table or step title bar across ``width`` columns."""
        self.merge(row, column, width, text, style)

    def note_box(
        self, row: int, column: int, width: int, height: int, text: str, style: str
    ) -> None:
        """A merged note with a medium outer border (the ⚠ boxes beside the HAG blocks)."""
        self.put(row, column, text, style)
        last_row, last_column = row + height - 1, column + width - 1
        for r in range(row, last_row + 1):
            for c in range(column, last_column + 1):
                if (r, c) == (row, column):
                    continue
                cell = self._claim(r, c)
                edge = _box_edge(r == row, r == last_row, c == column, c == last_column)
                if edge is not None:
                    self.styles.set(cell, edge)
        self.join(row, column, width, height)

    def notes(self, row: int, lines: Iterable[str], column: int = 1) -> int:
        """Small-print notes below a sheet, one per row; returns the next free row."""
        for line in lines:
            self.put(row, column, line, "note")
            row += 1
        return row

    def link(
        self, row: int, column: int, text: str, target_sheet: str, style: str = "link"
    ) -> None:
        """A cell that jumps to another sheet of the workbook."""
        self.put(row, column, text, style)
        cell = self.ws.cell(row, column)
        cell.hyperlink = Hyperlink(
            ref=cell.coordinate,
            location=f"{quote_sheetname(target_sheet)}!A1",
            display=text,
            tooltip=f"Go to {target_sheet}",
        )

    # ------------------------------------------------------------------ layout

    def widths(self, widths: dict[int | tuple[int, int], float]) -> None:
        """Column widths in characters: ``{1: 30, (2, 11): 11}`` (a tuple is a range of columns)."""
        for key, width in widths.items():
            first, last = (key, key) if isinstance(key, int) else key
            for column in range(first, last + 1):
                self.ws.column_dimensions[column_letter(column)].width = width + WIDTH_PADDING

    def hide_columns(self, first: int, last: int) -> None:
        """Hide helper columns."""
        for column in range(first, last + 1):
            self.ws.column_dimensions[column_letter(column)].hidden = True

    def height(self, row: int, height: float) -> None:
        """Row height in points."""
        self.ws.row_dimensions[row].height = height

    # ------------------------------------------------------------------ conditional highlights

    def highlight_equal(self, ref: str, value: str, kind: str = "good") -> None:
        """Highlight the cells of ``ref`` equal to ``value`` (``1`` or ``"="``)."""
        self.ws.conditional_formatting.add(
            ref,
            CellIsRule(operator="equal", formula=[value], **_differential(kind)),  # type: ignore[no-untyped-call]
        )

    def highlight_formula(self, ref: str, formula: str, kind: str) -> None:
        """Highlight the cells of ``ref`` where ``formula`` (written for its first cell) holds."""
        self.ws.conditional_formatting.add(
            ref,
            FormulaRule(formula=[formula], **_differential(kind)),  # type: ignore[no-untyped-call]
        )


def _pick(style: StyleChoice, index: int, row: int, column: int) -> str:
    """The style of one cell: ``index`` selects from a sequence, (row, column) feed a function."""
    if isinstance(style, str):
        return style
    if callable(style):
        return style(row, column)
    return style[index]


def _box_edge(top: bool, bottom: bool, left: bool, right: bool) -> str | None:
    if top:
        return "box_top_right" if right else "box_top"
    if bottom:
        if left:
            return "box_bottom_left"
        return "box_bottom_right" if right else "box_bottom"
    if left:
        return "box_left"
    return "box_right" if right else None


def _differential(kind: str) -> dict[str, Any]:
    fill, ink = HIGHLIGHTS[kind]
    return {
        "font": Font(bold=True, color="FF" + ink),
        "fill": PatternFill(bgColor="FF" + fill),
    }
