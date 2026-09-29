"""Data, metric, neighbours and synthetic-feature layer (formulas (1)–(6) of the article)."""
import openpyxl

ROUND = 10
K_MIN = 3

def load_dataset(path, sheet='Dataset'):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    X = [[float(ws.cell(r, c).value) for c in range(2, 15)] for r in range(3, 13)]
    y = [int(ws.cell(r, 15).value) for r in range(3, 13)]
    types = [int(ws.cell(13, c).value) for c in range(2, 15)]
    return X, y, types

# ---------------------------------------------------------------- operators (feature subsets of the Zhuravlev metric)
OPERATORS = [('Z', 'ρ — all 13 features'), ('I', 'ρ_I — quantitative features only'), ('J', 'ρ_J — nominal features only')]

def scale_params(Xtr, types):
    n = len(types)
    lo = [min(r[j] for r in Xtr) for j in range(n)]; hi = [max(r[j] for r in Xtr) for j in range(n)]
    return lo, hi

def unify(x, types, lo, hi):
    return [((x[j] - lo[j]) / (hi[j] - lo[j]) if hi[j] > lo[j] else 0.0) if types[j] == 1 else x[j] for j in range(len(types))]

def rho(zx, zy, x, y_, types, op='Z'):
    q = round(sum(abs(zx[j] - zy[j]) for j in range(len(types)) if types[j] == 1), ROUND) if op in ('Z', 'I') else 0.0
    nm = sum(1 for j in range(len(types)) if types[j] == 0 and x[j] != y_[j]) if op in ('Z', 'J') else 0
    return round(q + nm, ROUND)

def neighbours(dists, exclude=None):
    """indices sorted by (distance, original index); `exclude` = own index (not its own neighbour)."""
    return sorted((j for j in range(len(dists)) if j != exclude), key=lambda j: (dists[j], j))

def permitted_k(y):
    k_max = 2 * min(y.count(1), y.count(2)) - 3
    return list(range(K_MIN, k_max + 1, 2)), k_max

def chi1(order, ytr, k):                       # number of K1 objects among the k nearest (class-free for the object itself)
    return sum(1 for j in order[:k] if ytr[j] == 1)

def binary_a(c1, k):                           # formula (5): 1 if χ1 > [k/2], 2 if χ2 > [k/2]
    return 1 if c1 > k // 2 else 2

def membership(mu, y, k):
    """formula (1): f_k(μ) for μ = 0..k; None where no object has gradation μ."""
    n1, n2 = y.count(1), y.count(2)
    f = {}
    for g in range(k + 1):
        d1 = sum(1 for v, c in zip(mu, y) if v == g and c == 1); d2 = sum(1 for v, c in zip(mu, y) if v == g and c == 2)
        f[g] = dict(d1=d1, d2=d2, f=(d1 / n1) / (d1 / n1 + d2 / n2) if d1 + d2 else None)
    return f

def stability(fk, m, beta):
    """formula (2): g = 1/m Σ n(μ)·f if f ≥ 0.5 else n(μ)·(1 − f)."""
    s = 0.0
    for g in range(beta + 1):
        e = fk.get(g)
        if e and e['f'] is not None:
            n_ = e['d1'] + e['d2']; s += n_ * (e['f'] if e['f'] >= 0.5 else 1 - e['f'])
    return s / m

def boundary(fk):
    """formula (3): G = (q1 + q2)/2, q2 = max{f < 0.5}, q1 = min{f > 0.5}; one side missing → 0.5."""
    lo = [e['f'] for e in fk.values() if e['f'] is not None and e['f'] < 0.5]
    hi = [e['f'] for e in fk.values() if e['f'] is not None and e['f'] > 0.5]
    q2 = max(lo) if lo else None; q1 = min(hi) if hi else None
    G = (q1 + q2) / 2 if (lo and hi) else 0.5
    return G, q1, q2

def informativeness(mu, y, fk, G):
    """formula (4): ω = (|{S∈K1 : f(μ_S) > G}| + |{S∈K2 : f(μ_S) < G}|) / m."""
    ok = sum(1 for v, c in zip(mu, y) if (c == 1 and fk[v]['f'] > G) or (c == 2 and fk[v]['f'] < G))
    return ok / len(y)

def contributions(a, y, w):
    """formula (6): η(j) = ω (α¹_j/|K1| − α²_j/|K2|), j ∈ {1, 2}."""
    n1, n2 = y.count(1), y.count(2); out = {}
    for j in (1, 2):
        a1 = sum(1 for v, c in zip(a, y) if v == j and c == 1); a2 = sum(1 for v, c in zip(a, y) if v == j and c == 2)
        out[j] = dict(a1=a1, a2=a2, eta=w * (a1 / n1 - a2 / n2))
    return out
