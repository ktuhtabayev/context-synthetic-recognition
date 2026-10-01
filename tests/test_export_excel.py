"""The Excel mirror: the writer, the palette, the layout on every shape and the size limits."""

import json
import warnings
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import numpy as np
import openpyxl
import pytest
from openpyxl.workbook import Workbook

from context_synthetic_recognition.config.models import CentreMode
from context_synthetic_recognition.export import RunView
from context_synthetic_recognition.export.excel import ExcelOptions, build_workbook, write_workbook
from context_synthetic_recognition.export.excel.common import (
    DISTANCES,
    DISTANCES_OTHER,
    blocks_that_fit,
    greedy_sheets,
    plan_sheets,
)
from context_synthetic_recognition.export.excel.styles import STYLES, TAB_COLORS, StyleBook
from context_synthetic_recognition.export.excel.template_data import (
    template_hag,
    template_input,
    template_variants,
)
from context_synthetic_recognition.export.excel.writer import (
    LayoutError,
    SheetWriter,
    cell_value,
    column_letter,
    reference,
    sheet_name,
)
from context_synthetic_recognition.services.validation.templates import read_hag_template

from .conftest import HAG_TEMPLATE
from .shapes import SHAPES, shape

EXPERIMENT_SHEETS = [
    "Overview",
    "Parameters",
    "Template Deviations",
    "Dataset",
    "Quantitative",
    "Nominal",
    "Normalized Dataset",
    "Zhuravlyov Distances",
    "Sorted Neighbors (ρ)",
    "Sorted Neighbors (ρ_I)",
    "Sorted Neighbors (ρ_J)",
    "Synthetic Features (k-NN)",
    "Ψ(r) Binary Features",
    "Membership & Stability",
    "Informativeness ω",
    "Ψ(r) Contribution & Weight",
    "Greedy upon Weight (1-Latent)",
    "Greedy upon Weight (2-Latent)",
    "Greedy upon Weight (3-Latent)",
    "Greedy upon Weight (4-Latent)",
    "Dataset for Meta-algorithm",
    "Brace for Meta-algorithm",
    "Meta-algorithm",
    "Meta-algorithm (All Objects)",
    "Margin Analysis",
    "Accuracy",
    "Confusion Matrix",
    "Precision, Recall, F1 Score",
    "ROC Curve & AUC",
    "Leave-One-Out",
    "Sensitivity (Switches)",
    "Model Properties",
    "Validation",
]
"""The sheets of the Excel experiment, in its order (the metric's name spelled Zhuravlyov)."""

SMALL_BUDGET = ExcelOptions(check_overlaps=True, max_sheet_cells=4000, max_matrix_cells=8000)


# ---------------------------------------------------------------- the writer


def test_cell_values() -> None:
    assert cell_value(None) is None
    assert cell_value("ρ") == "ρ"
    assert cell_value(True) == 1
    assert cell_value(np.bool_(False)) == 0
    assert cell_value(np.int64(7)) == 7
    assert type(cell_value(np.int64(7))) is int
    assert cell_value(np.float64(0.25)) == 0.25
    assert type(cell_value(np.float64(0.25))) is float
    assert cell_value(float("nan")) == "—"
    assert cell_value(float("inf")) == "+∞"
    assert cell_value(float("-inf")) == "−∞"


def test_references_and_sheet_names() -> None:
    assert column_letter(1) == "A"
    assert column_letter(27) == "AA"
    assert reference(5, 2) == "B5"
    assert reference(5, 2, rows=10, columns=6) == "B5:G14"
    assert sheet_name("Sorted Neighbors (ρ_I)") == "Sorted Neighbors (ρ_I)"
    assert sheet_name("a/b:c*d?[e]\\f") == "a·b·c·d··e··f"
    assert sheet_name("'quoted'") == "quoted"
    assert sheet_name("") == "Sheet"
    long = sheet_name("Sorted Neighbors (a very long operator label)")
    assert len(long) == 31
    second = sheet_name("Sorted Neighbors (a very long operator label)", [long])
    assert second.endswith(" (2)")
    assert len(second) == 31
    assert sheet_name("Dataset", ["dataset", "Dataset (2)"]) == "Dataset (3)"  # case-blind


def writer(check_overlaps: bool = False) -> SheetWriter:
    workbook = Workbook()
    return SheetWriter(
        workbook, StyleBook(workbook), "Test", "hag", freeze="B4", check_overlaps=check_overlaps
    )


def test_the_sheet_set_up() -> None:
    sheet = writer()
    ws = sheet.ws
    assert sheet.title == "Test"
    assert ws.sheet_properties.tabColor.rgb == "FF" + TAB_COLORS["hag"]
    assert ws.freeze_panes == "B4"
    assert ws.sheet_view.showGridLines is False
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.oddFooter.center.text == "Page &P of &N"


def test_writing_cells() -> None:
    sheet = writer()
    ws = sheet.ws
    sheet.put(1, 1, np.float64(0.5), "num")
    assert ws["A1"].value == 0.5
    assert ws["A1"].number_format == "0.0000"
    sheet.row(2, 2, ["a", "b", "c"], "header")
    assert [ws.cell(2, c).value for c in (2, 3, 4)] == ["a", "b", "c"]
    assert ws["C2"].font.bold
    sheet.row(3, 1, ["S₁", 1.5], ("label", "num"))
    assert ws["A3"].fill.fgColor.rgb != ws["B3"].fill.fgColor.rgb
    sheet.column(4, 1, [1, 2, 3], lambda row, column: "integer" if row else "header")
    assert ws["A4"].font.bold
    assert not ws["A5"].font.bold
    sheet.block(8, 1, np.array([[1.0, 2.0], [float("nan"), 4.0]]), ("num", "integer"))
    assert ws["A9"].value == "—"
    assert ws["B9"].value == 4.0
    assert ws["A8"].number_format == "0.0000"
    assert ws["B8"].number_format != "0.0000"
    assert sheet.cells == 1 + 3 + 2 + 3 + 4
    with pytest.raises(KeyError):
        sheet.put(20, 1, 1, "no such style")


def test_merging() -> None:
    sheet = writer()
    ws = sheet.ws
    sheet.title_bar("Title", 6)
    assert "A1:F1" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["A1"].value == "Title"
    assert ws.row_dimensions[1].height == 20.0
    # the covered cells carry the border of the bar: openpyxl would only do that on merge_cells
    assert ws["F1"].border.right.style == "thin"
    assert ws["C1"].border.top.style == "thin"
    sheet.bar(3, 2, 4, "STEP 1")
    assert "B3:E3" in [str(r) for r in ws.merged_cells.ranges]
    sheet.merge(5, 1, 1, "single", "label")
    assert len(ws.merged_cells.ranges) == 2
    sheet.merge(6, 1, 2, "tall", "label", height=3)
    assert "A6:B8" in [str(r) for r in ws.merged_cells.ranges]
    sheet.note_box(10, 1, 3, 3, "⚠ note", "warn1_box")
    assert "A10:C12" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["A10"].value == "⚠ note"
    assert ws["C10"].border.right.style == "medium"
    assert ws["A12"].border.bottom.style == "medium"
    assert ws["B12"].border.bottom.style == "medium"
    assert ws["C12"].border.right.style == "medium"
    assert ws["A11"].border.left.style == "medium"
    assert ws["C11"].border.right.style == "medium"
    assert ws["B11"].border.left.style is None  # the middle of the box has no edge


def test_notes_links_and_dimensions() -> None:
    sheet = writer()
    ws = sheet.ws
    assert sheet.notes(4, ["first", "second"]) == 6
    assert ws["A5"].value == "second"
    sheet.link(8, 2, "Dataset", "Ψ(r) Binary Features")
    assert ws["B8"].hyperlink.location == "'Ψ(r) Binary Features'!A1"
    assert ws["B8"].hyperlink.tooltip == "Go to Ψ(r) Binary Features"
    sheet.widths({1: 30, (2, 4): 11})
    assert ws.column_dimensions["A"].width == pytest.approx(30.77734375)
    assert ws.column_dimensions["D"].width == pytest.approx(11.77734375)
    sheet.hide_columns(6, 7)
    assert ws.column_dimensions["F"].hidden
    assert ws.column_dimensions["G"].hidden
    sheet.height(3, 31.5)
    assert ws.row_dimensions[3].height == 31.5
    sheet.highlight_equal("B2:D9", "1")
    sheet.highlight_formula("B10:D12", "B10=MIN($B10:$D10)", "bad")
    rules = {
        str(rng.sqref): [(rule.type, rule.operator, tuple(rule.formula)) for rule in found]
        for rng, found in ws.conditional_formatting._cf_rules.items()
    }
    assert rules == {
        "B2:D9": [("cellIs", "equal", ("1",))],
        "B10:D12": [("expression", None, ("B10=MIN($B10:$D10)",))],
    }


def test_overlapping_tables_are_detected() -> None:
    sheet = writer(check_overlaps=True)
    sheet.block(1, 1, [[1, 2], [3, 4]], "integer")
    with pytest.raises(LayoutError, match=r"Test!B2 is written twice"):
        sheet.put(2, 2, 5, "integer")
    sheet.merge(5, 1, 3, "bar", "bar")
    with pytest.raises(LayoutError, match=r"Test!C5"):
        sheet.put(5, 3, "inside the bar", "text")
    # without the check a later write wins silently
    free = writer()
    free.put(1, 1, 1, "integer")
    free.put(1, 1, 2, "integer")
    assert free.ws["A1"].value == 2


def test_the_palette() -> None:
    workbook = Workbook()
    styles = StyleBook(workbook)
    assert workbook.sheetnames == ["Sheet"]  # the scratch sheet is removed again
    ws = workbook.active
    for column, name in enumerate(STYLES, start=1):
        styles.set(ws.cell(1, column), name)
    by_name = {name: ws.cell(1, column) for column, name in enumerate(STYLES, start=1)}
    title = by_name["title"]
    assert title.font.bold
    assert title.font.color.rgb == "FFFFFFFF"
    assert title.fill.fgColor.rgb == "FF1F3864"
    assert by_name["num"].number_format == "0.0000"
    assert by_name["yes_no"].number_format == '"Yes";"Yes";"No"'
    assert by_name["header_k"].number_format == '"k = "0'
    assert by_name["text_wrap"].alignment.wrap_text
    assert by_name["quantitative"].fill.fgColor.rgb == "FFDDEBF7"
    assert by_name["nominal"].fill.fgColor.rgb == "FFFFF2CC"
    assert all(cell.font.name == "Calibri" for cell in by_name.values())


def test_blocks_that_fit() -> None:
    assert blocks_that_fit(10, 100, 1000) == 10
    assert blocks_that_fit(10, 100, 950) == 9
    assert blocks_that_fit(10, 100, 50) == 1  # at least one block is always written
    assert blocks_that_fit(0, 100, 50) == 0


# ---------------------------------------------------------------- the mirror of the experiment


@pytest.fixture(scope="module")
def mirror(experiment_view: RunView) -> Any:
    return build_workbook(experiment_view, ExcelOptions(check_overlaps=True))


def test_the_mirror_has_the_sheets_of_the_experiment(mirror: Any, experiment_view: RunView) -> None:
    workbook, book = mirror
    assert workbook.sheetnames == EXPERIMENT_SHEETS
    assert book.omitted == []
    assert workbook.active.title == "Overview"
    assert [ws.title for ws in workbook.worksheets if ws.sheet_view.tabSelected] == ["Overview"]
    assert workbook.properties.title == f"Context-Synthetic Model — {experiment_view.dataset.name}"
    assert workbook.properties.creator.startswith("context-synthetic-recognition ")
    assert book.cells == sum(sheet.cells for sheet in book.sheets) > 15_000
    for ws in workbook.worksheets:
        title = ws["B2"] if ws.title == "Overview" else ws["A1"]
        assert title.value, ws.title
        assert title.font.bold, ws.title
        assert ws.sheet_properties.tabColor is not None, ws.title
        assert ws.sheet_view.showGridLines is False, ws.title


def test_the_mirror_holds_values_not_formulas(mirror: Any) -> None:
    workbook, _ = mirror
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                # the HAG sheets mark equal class centres with the text "=": not a formula
                assert cell.data_type != "f", f"{ws.title}!{cell.coordinate}"


def test_the_overview_links_to_every_sheet(mirror: Any) -> None:
    workbook, _ = mirror
    overview = workbook["Overview"]
    targets = [
        cell.hyperlink.location for row in overview.iter_rows() for cell in row if cell.hyperlink
    ]
    # the list of sheets and the stages of the pipeline both link to the sheets
    assert set(targets) == {f"'{name}'!A1" for name in EXPERIMENT_SHEETS[1:]}


def test_the_template_deviations_are_highlighted(mirror: Any) -> None:
    workbook, _ = mirror
    parameters = [
        str(cell.value)
        for row in workbook["Parameters"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert "⚠1 Class centres in θ and γ" in parameters
    assert "⚠2 Majorizer passes in STEP 4" in parameters
    assert str(workbook["Validation"]["A3"].value).startswith(
        "⚠ Template calculations that differ from the article"
    )
    deviations = workbook["Template Deviations"]
    assert str(deviations["A1"].value).startswith("Template Deviations — where the template HAG")
    text = [str(c.value) for row in deviations.iter_rows() for c in row if c.value is not None]
    assert "{x₃, x₆, x₁₃, x₄, x₉}" in text  # the template data under the template calculation
    assert "{x₃, x₁, x₅, x₉, x₆}" in text  # … and under the article's
    assert "< 1e-9  (identical)" in text


def test_the_mirror_spells_zhuravlyov(mirror: Any) -> None:
    workbook, _ = mirror
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert "Zhuravlev" not in str(cell.value), f"{ws.title}!{cell.coordinate}"


def test_the_mirror_is_written_and_read_back(experiment_view: RunView, tmp_path: Path) -> None:
    path = tmp_path / "deep" / "workbook.xlsx"
    book = write_workbook(experiment_view, path)
    assert book.omitted == []
    loaded = openpyxl.load_workbook(path)
    assert loaded.sheetnames == EXPERIMENT_SHEETS
    assert loaded.properties.created == experiment_view.created.replace(  # type: ignore[union-attr]
        tzinfo=None, microsecond=0
    )
    loaded.close()


# ---------------------------------------------------------------- other shapes


def expected_sheets(view: RunView) -> list[str]:
    """The sheets a run needs: the experiment's, with the parts that depend on the run."""
    trace = view.trace
    quantitative = bool(trace.scaling.quantitative.any())
    nominal = bool((~trace.scaling.quantitative).any())
    sheets = ["Overview", "Parameters", "Template Deviations", "Dataset"]
    sheets += ["Quantitative"] * quantitative + ["Nominal"] * nominal
    sheets += ["Normalized Dataset", "Zhuravlyov Distances"]
    sheets += [f"Sorted Neighbors ({operator.label})" for operator in trace.operators]
    sheets += [
        "Synthetic Features (k-NN)",
        "Ψ(r) Binary Features",
        "Membership & Stability",
        "Informativeness ω",
        "Ψ(r) Contribution & Weight",
    ]
    sheets += [f"Greedy upon Weight ({j}-Latent)" for j in range(1, greedy_sheets(view) + 1)]
    sheets += [
        "Dataset for Meta-algorithm",
        "Brace for Meta-algorithm",
        "Meta-algorithm",
        "Meta-algorithm (All Objects)",
        "Margin Analysis",
        "Accuracy",
        "Confusion Matrix",
        "Precision, Recall, F1 Score",
        "ROC Curve & AUC",
    ]
    titles = {
        "leave-one-out": "Leave-One-Out",
        "stratified-k-fold": "Stratified K-Fold",
        "repeated-k-fold": "Repeated K-Fold",
        "hold-out": "Hold-Out",
    }
    sheets += [titles[p.protocol] for p in view.cross_validations]
    sheets += ["Sensitivity (Switches)"] * (view.sensitivity is not None)
    return [*sheets, "Model Properties", "Validation"]


@pytest.mark.parametrize("name", sorted(SHAPES))
def test_the_layout_of_other_shapes_does_not_overlap(name: str) -> None:
    view = shape(name)
    workbook, book = build_workbook(view, SMALL_BUDGET)  # LayoutError if two tables overlap
    assert workbook.sheetnames == expected_sheets(view)
    for message in book.omitted:
        sheet = message.split(":")[0]
        assert sheet in workbook.sheetnames, message
        # what was left out is also said on the sheet itself
        text = [
            str(c.value) for row in workbook[sheet].iter_rows() for c in row if c.value is not None
        ]
        assert any(message.split(": ", 1)[1][:40] in t for t in text), message


def test_the_number_of_sheets_follows_the_run() -> None:
    numeric, nominal, tiny = shape("numeric"), shape("nominal"), shape("tiny")
    assert numeric.slots == 3
    assert greedy_sheets(numeric) == 2  # ϰ = 3
    assert greedy_sheets(nominal) == 0  # r = 1: nothing to group
    assert nominal.hag.p == 0
    assert greedy_sheets(tiny) == 1
    plan = plan_sheets(numeric)
    assert plan.quantitative
    assert not plan.nominal
    assert not plan.sensitivity
    assert plan.neighbours == ("Sorted Neighbors (ρ)", "Sorted Neighbors (ρ_I)")
    assert plan.protocols == {"leave-one-out": "Leave-One-Out", "hold-out": "Hold-Out"}
    assert plan_sheets(nominal).protocols == {"repeated-k-fold": "Repeated K-Fold"}


def test_the_distance_sheet_is_named_after_the_metric() -> None:
    def view(metric: str, protocol: str) -> Any:
        return SimpleNamespace(
            trace=SimpleNamespace(
                operators=[SimpleNamespace(label="ρ", metric=metric)],
                scaling=SimpleNamespace(quantitative=np.array([True, False])),
            ),
            slots=2,
            hag=SimpleNamespace(iterations=[None]),
            cross_validations=[SimpleNamespace(protocol=protocol)],
            sensitivity=None,
        )

    assert plan_sheets(view("zhuravlyov", "leave-one-out")).distances == DISTANCES
    other = plan_sheets(view("euclidean", "monte-carlo"))
    assert other.distances == DISTANCES_OTHER == "Distances"
    assert other.protocols == {"monte-carlo": "Monte-Carlo"}  # an unknown protocol: its title


def test_the_size_limits_leave_out_and_say_so() -> None:
    _, book = build_workbook(shape("heart270"), SMALL_BUDGET)
    omitted = "\n".join(book.omitted)
    assert "the 3 matrices of 270 × 270 distances are not written" in omitted
    assert "Sorted Neighbors (ρ_J): the 270 × 270 rank matrix is not written" in omitted
    assert "ranks 1 … 22 (the largest permitted k and one beyond) for 22 of 270 objects" in omitted
    assert "only the block of the chosen feature q is written (30 candidates)" in omitted
    assert "B1/B2 flags per step are written for 1 of 270 objects" in omitted
    assert "Task 2 (bit representations) is not written" in omitted
    _, literal = build_workbook(shape("literal"), SMALL_BUDGET)
    omitted = "\n".join(literal.omitted)
    assert "those of the TUPLAM features only (5 of 354)" in omitted
    assert "the 354 × 354 table of identical synthetic features is not written" in omitted
    # with the default limits a small run is written in full
    _, full = build_workbook(shape("numeric"), ExcelOptions(check_overlaps=True))
    assert full.omitted == []
    _, tight = build_workbook(shape("numeric"), SMALL_BUDGET)
    assert tight.cells < full.cells
    assert any("only the block of a₇ is written" in m for m in tight.omitted)  # not executed


@pytest.mark.parametrize("name", ["numeric", "nominal", "tiny"])
def test_other_shapes_are_written_and_read_back(name: str, tmp_path: Path) -> None:
    view = shape(name)
    path = tmp_path / f"{name}.xlsx"
    write_workbook(view, path, ExcelOptions(check_overlaps=True))
    loaded = openpyxl.load_workbook(path)
    assert loaded.sheetnames == expected_sheets(view)
    dataset = loaded["Dataset"]
    values = [c.value for row in dataset.iter_rows() for c in row if c.value is not None]
    assert view.trace.classes[0] in values  # text labels are written as they are
    loaded.close()


# ---------------------------------------------------------------- the template data


def test_the_template_data_is_the_template_workbook() -> None:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = openpyxl.load_workbook(HAG_TEMPLATE, data_only=True)
    contributions, weights, class_index, labels = read_hag_template(book)
    sheet = book["Dataset for Meta-algorithm"]
    latent = [[sheet.cell(r, c).value for c in range(10, 14)] for r in range(3, 13)]
    book.close()
    data = template_input()
    assert np.array_equal(data.contributions, contributions)
    assert np.array_equal(data.weights, weights)
    assert np.array_equal(data.class_index, class_index)
    assert data.labels == tuple(labels)
    assert np.array_equal(data.latent, np.array(latent))


def test_the_template_resource_is_small_json() -> None:
    from importlib.resources import files

    resource = files("context_synthetic_recognition.export.excel") / "hag_template.json"
    data = json.loads(resource.read_text(encoding="utf-8"))
    assert set(data) == {"source", "contributions", "weights", "classes", "latent"}
    assert (len(data["contributions"]), len(data["contributions"][0])) == (10, 13)


def test_the_switch_settings_on_the_template_data() -> None:
    variants = template_variants()
    assert [(v.centres, v.step4_passes, v.label) for v in variants] == [
        (CentreMode.RUNNING, 2, "template cells"),
        (CentreMode.RUNNING, 1, "running centres, 1 pass"),
        (CentreMode.FINAL, 2, "final centres, 2 passes"),
        (CentreMode.FINAL, 1, "article"),
    ]
    assert [v.selected for v in variants] == [
        "{x₃, x₆, x₁₃, x₄, x₉}",
        "{x₃, x₆, x₁₃, x₅, x₁₂}",
        "{x₃, x₁, x₅, x₉}",
        "{x₃, x₁, x₅, x₉, x₆}",
    ]
    # the template setting reproduces the template's latent features; the others do not
    assert variants[0].difference < 1e-12
    assert variants[0].difference_text == "< 1e-9  (identical)"
    assert [v.difference_text for v in variants[1:]] == ["1.2e+00", "1.2e+00", "1.7e+00"]
    assert template_hag().p == 4
    assert template_hag(CentreMode.FINAL, 2).p == 3
