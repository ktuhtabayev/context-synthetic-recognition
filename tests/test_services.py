"""Services: workbook validation (failure paths) and dataset summaries."""

import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from context_synthetic_recognition.config import ExperimentConfig, HAGConfig, plugin, preset
from context_synthetic_recognition.data import Dataset, FeatureType, load_builtin
from context_synthetic_recognition.errors import DatasetError
from context_synthetic_recognition.services.datasets import parse_object, summarize
from context_synthetic_recognition.services.validation import (
    ValidationError,
    WorkbookKind,
    validate_workbook,
    workbook_kind,
)

from .conftest import HAG_TEMPLATE, META_TEMPLATE, WORKBOOK


def values_copy(
    tmp_path: Path, edits: dict[tuple[str, str], object], source: Path = WORKBOOK
) -> Path:
    """A copy of a workbook with cached values only (formulas dropped) and some cells changed."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = openpyxl.load_workbook(source, data_only=True)
    for (sheet, cell), value in edits.items():
        book[sheet][cell] = value
    path = tmp_path / "copy.xlsx"
    book.save(path)
    book.close()
    return path


def test_a_changed_cell_is_reported(tmp_path: Path) -> None:
    path = values_copy(
        tmp_path,
        {("Zhuravlev Distances", "C5"): 6.2, ("Ψ(r) Binary Features", "B40"): "x"},
    )
    report = validate_workbook(path)
    assert not report.passed
    failed = [r for r in report.results if not r.passed]
    assert [(r.sheet, r.ref) for r in failed] == [
        ("Zhuravlev Distances", "B5:K14"),
        ("Ψ(r) Binary Features", "B40:G49"),
    ]
    assert failed[0].mismatches == ("C5: workbook 6.2 ≠ package 6.1621313645",)
    assert failed[1].mismatches == ("B40: workbook 'x' ≠ package 2",)


def test_a_layout_that_does_not_fit_the_data(tmp_path: Path) -> None:
    # S₂ moved from K1 to K2: |K1| = 3, so only k = 3 is permitted
    path = values_copy(tmp_path, {("Dataset", "O4"): 2})
    with pytest.raises(ValidationError, match="1 permitted k"):
        validate_workbook(path)


def test_a_missing_sheet(tmp_path: Path) -> None:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = openpyxl.load_workbook(WORKBOOK, data_only=True)
    del book["Informativeness ω"]
    path = tmp_path / "short.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ValidationError, match="sheet 'Informativeness ω' is missing"):
        validate_workbook(path)


def test_the_hag_switches_are_read_from_the_parameters_sheet(tmp_path: Path) -> None:
    # the cached values stay those of the template setting, so the article setting must differ
    path = values_copy(tmp_path, {("Parameters", "B24"): 2, ("Parameters", "B25"): 1})
    report = validate_workbook(path)
    assert "centres = final, STEP 4 passes = 1" in report.notes[0]
    assert not any("ADR-00" in note for note in report.notes)
    failed = {r.sheet for r in report.results if not r.passed}
    assert "Greedy upon Weight (1-Latent)" in failed
    assert "Zhuravlev Distances" not in failed  # Steps 1–8 do not depend on the switches


def test_a_configuration_that_differs_from_the_workbook() -> None:
    report = validate_workbook(WORKBOOK, preset("article"))
    assert not report.passed
    assert report.notes[1].startswith("⚠ the configuration's HAG settings differ")


def test_a_genuinely_new_object(tmp_path: Path) -> None:
    # B35 = 0: S₁'s values classified as a new object — S₁ is then its own nearest neighbour
    report = validate_workbook(values_copy(tmp_path, {("Brace for Meta-algorithm", "B35"): 0}))
    failed = [r for r in report.results if not r.passed]
    # its ranks and χ₁ change; the cached cells hold the left-out version
    assert [r.ref for r in failed if r.sheet == "Brace for Meta-algorithm"] == [
        "B45:K47",
        "B53:G56",
    ]
    assert failed[0].mismatches[0] == "B45: workbook 999 ≠ package 1"


@pytest.mark.parametrize(
    ("edit", "message"),
    [
        (("Parameters", "B19"), r"Parameters!B19 \(α\)"),
        (("Parameters", "B24"), "the switches must be 1 or 2"),
        (("Brace for Meta-algorithm", "B35"), "no training object № 11"),
        (("Brace for Meta-algorithm", "E31"), "is not a number"),
    ],
)
def test_bad_workbook_inputs(tmp_path: Path, edit: tuple[str, str], message: str) -> None:
    value = {"B19": "x", "B24": 3, "B35": 11, "E31": None}[edit[1]]
    with pytest.raises(ValidationError, match=message):
        validate_workbook(values_copy(tmp_path, {edit: value}))


def test_shorter_groupings_are_mapped_too() -> None:
    # ϰ = 3: iterations 3 and 4 are not executed; the map then expects the workbook's
    # "iteration not executed" states and "—" for the unfilled TUPLAM positions
    report = validate_workbook(WORKBOOK, ExperimentConfig(hag=HAGConfig(kappa=3)))
    by_ref = {(r.sheet, r.ref): r for r in report.results}
    d4 = by_ref[("Greedy upon Weight (3-Latent)", "D4")]
    assert d4.mismatches == ("D4: workbook 1 ≠ package 0",)
    slots = by_ref[("Dataset for Meta-algorithm", "C5:C9")]
    assert slots.mismatches == (
        "C8: workbook 'a₂' ≠ package '—'",
        "C9: workbook 'a₄' ≠ package '—'",
    )
    assert by_ref[("Greedy upon Weight (1-Latent)", "D4")].passed
    # no candidate below cr1₀: iteration 1 runs but adds nothing
    stopped = validate_workbook(WORKBOOK, ExperimentConfig(hag=HAGConfig(cr1=0.4)))
    q = {(r.sheet, r.ref): r for r in stopped.results}[
        ("Greedy upon Weight (1-Latent)", "J112:K112")
    ]
    assert q.mismatches == ("J112: workbook 3 ≠ package '—'", "K112: workbook 'a₃' ≠ package '—'")


def test_a_hag_longer_than_the_layout(tmp_path: Path) -> None:
    with pytest.raises(ValidationError, match="more than 4 iterations"):
        validate_workbook(values_copy(tmp_path, {("Parameters", "B21"): 6}))


def test_workbook_kinds(tmp_path: Path) -> None:
    assert workbook_kind(["Parameters", "Dataset"]) is WorkbookKind.EXPERIMENT
    assert (
        workbook_kind(["Dataset (Contribution & Weight)", "Greedy upon Weight (1-Latent)"])
        is WorkbookKind.HAG_TEMPLATE
    )
    assert (
        workbook_kind(["Brace for Meta-algorithm", "Meta-algorithm"]) is WorkbookKind.META_TEMPLATE
    )
    book = openpyxl.Workbook()
    path = tmp_path / "other.xlsx"
    book.save(path)
    with pytest.raises(ValidationError, match="not a known workbook"):
        validate_workbook(path)
    with pytest.raises(ValidationError, match="cannot read the workbook"):
        validate_workbook(tmp_path / "missing.xlsx")


def test_template_inputs_are_checked(tmp_path: Path) -> None:
    hag = values_copy(tmp_path, {("Dataset (Contribution & Weight)", "O3"): 3}, HAG_TEMPLATE)
    with pytest.raises(ValidationError, match="classes must be 1 or 2"):
        validate_workbook(hag)
    meta = values_copy(tmp_path, {("Brace for Meta-algorithm", "K88"): 0}, META_TEMPLATE)
    with pytest.raises(ValidationError, match="classes must be 1 or 2"):
        validate_workbook(meta)
    changed = values_copy(tmp_path, {("Brace for Meta-algorithm", "B77"): 3}, META_TEMPLATE)
    report = validate_workbook(changed)
    assert not report.passed  # the new object changed, the cached result did not


def test_parse_object(heart270: Dataset) -> None:
    values = parse_object(heart270, "70; 1 4,130 322 0 2 109 0 2.4 2 3 3")
    assert values.tolist()[:3] == [70.0, 1.0, 4.0]
    with pytest.raises(DatasetError, match="expected 13 values"):
        parse_object(heart270, "1 2")
    with pytest.raises(DatasetError, match="x₂: 'male' is not a number"):
        parse_object(heart270, "70 male 4 130 322 0 2 109 0 2.4 2 3 3")
    labelled = Dataset(
        X=np.array([[0.0, 1.0], [1.0, 0.0]]),
        y=np.array([1, 2]),
        feature_types=(FeatureType.NOMINAL, FeatureType.QUANTITATIVE),
        feature_names=("sex", "age"),
        categories={"sex": ("female", "male")},
    )
    assert parse_object(labelled, "male 0.5").tolist() == [1.0, 0.5]
    with pytest.raises(DatasetError, match="categories: female, male"):
        parse_object(labelled, "other 0.5")


# ---------------------------------------------------------------- dataset summaries


def test_summary_of_heart_disease_270() -> None:
    summary = summarize(load_builtin("heart-disease-270"))
    assert summary.permitted_k is not None
    assert summary.permitted_k.count == 118
    assert summary.r == 354
    assert [op.label for op in summary.operators] == ["ρ", "ρ_I", "ρ_J"]
    assert summary.operators[2].features == ("x₂", "x₃", "x₆", "x₇", "x₉", "x₁₁", "x₁₃")
    assert summary.problems == ()


def test_summary_reports_why_the_model_is_undefined() -> None:
    three = Dataset(
        X=[[1.0], [2.0], [3.0], [4.0]], y=[1, 2, 3, 3], feature_types=(FeatureType.QUANTITATIVE,)
    )
    summary = summarize(three)
    assert summary.permitted_k is None
    assert summary.r is None
    assert any("binary" in p for p in summary.problems)
    assert any("no k is permitted" in p for p in summary.problems)
    assert [s.label for s in summary.skipped_operators] == ["ρ_J"]


def test_summary_with_a_capped_k_range() -> None:
    config = ExperimentConfig(k=plugin("formula", k_max_cap=7))
    summary = summarize(load_builtin("heart-disease-270"), config)
    assert summary.permitted_k is not None
    assert summary.permitted_k.ks == (3, 5, 7)
    assert summary.r == 9
