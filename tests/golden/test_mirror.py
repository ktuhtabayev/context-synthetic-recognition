"""Golden tests of the Excel mirror: the export of the experiment against the experiment workbook.

Two acceptance checks (ADR-037):

* the exported workbook passes ``csr validate`` — every one of the 11,631 computed cells of the
  validation map sits at its original address with its value (tolerance 1e-9);
* sheet by sheet it has the structure and the style of the original: the same sheets in the same
  order, every number at the same address, the same merged ranges, frozen panes, tab colours,
  conditional formats, column widths and cell styles. What differs is listed here and is
  intended (the mirror holds values, spells Zhuravlyov, and adds the run's identification).
"""

import warnings
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from openpyxl.cell import Cell

from context_synthetic_recognition.export import RunView
from context_synthetic_recognition.export.excel import ExcelOptions, write_workbook
from context_synthetic_recognition.services.validation import validate_workbook

from ..conftest import WORKBOOK

pytestmark = pytest.mark.golden

RENAMED = {"Zhuravlev Distances": "Zhuravlyov Distances"}
"""Original sheet name → the mirror's (the transliteration of the article, ADR-039)."""
RUN_BLOCK = ("Parameters", range(39, 48))
"""Rows the mirror adds below the original content: "This run · what identifies it"."""

TEXT_DIFFERENCES = {
    # the original shows these labels as Office-Math drawings, the mirror writes them as text
    "Dataset": 56,
    "Quantitative": 16,
    "Nominal": 17,
    # wording adapted from "live formulas / the engine" to a static export, and Zhuravlyov
    "Overview": 34,
    "Parameters": 3,
    "Template Deviations": 5,
    "Zhuravlev Distances": 4,
    "Greedy upon Weight (1-Latent)": 1,
    "Greedy upon Weight (2-Latent)": 1,
    "Greedy upon Weight (3-Latent)": 1,
    "Greedy upon Weight (4-Latent)": 1,
    "Brace for Meta-algorithm": 2,
    "Accuracy": 8,
    "Confusion Matrix": 2,
    "Precision, Recall, F1 Score": 24,
    "ROC Curve & AUC": 2,
    "Leave-One-Out": 3,
    "Sensitivity (Switches)": 4,
    "Validation": 15,
}
"""Sheet → the largest number of text cells that may differ from the original."""

STYLE_DIFFERENCES = {
    # quirks of the original the mirror does not copy: a missing border, two stray bold cells
    "Dataset": {"R2", "Y2", "H20"},
    "Quantitative": {"H2", "I15"},
}
WIDTH_DIFFERENCES = {
    "Dataset": {"C"},
    "Quantitative": {"K", "L", "Q"},
    "Nominal": {"I"},
}
"""Columns whose width follows the generic layout rather than the original's hand-set value."""


@pytest.fixture(scope="module")
def original() -> Iterator[Any]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = openpyxl.load_workbook(WORKBOOK, data_only=True)
    yield book
    book.close()


@pytest.fixture(scope="module")
def mirror_path(experiment_view: RunView, tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("mirror") / "workbook.xlsx"
    book = write_workbook(experiment_view, path, ExcelOptions(check_overlaps=True))
    assert book.omitted == []
    return path


@pytest.fixture(scope="module")
def mirror(mirror_path: Path) -> Iterator[Any]:
    book = openpyxl.load_workbook(mirror_path)
    yield book
    book.close()


def pairs(original: Any, mirror: Any) -> Iterator[tuple[str, Any, Any]]:
    for name in original.sheetnames:
        yield name, original[name], mirror[RENAMED.get(name, name)]


def is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def empty(value: Any) -> bool:
    return value is None or value == ""


def cells(a: Any, b: Any) -> Iterator[tuple[Any, Any]]:
    for row in range(1, max(a.max_row, b.max_row) + 1):
        for column in range(1, max(a.max_column, b.max_column) + 1):
            yield a.cell(row, column), b.cell(row, column)


def in_run_block(sheet: str, cell: Any) -> bool:
    return sheet == RUN_BLOCK[0] and cell.row in RUN_BLOCK[1]


# ---------------------------------------------------------------- csr validate


def test_the_mirror_passes_the_validation(mirror_path: Path) -> None:
    report = validate_workbook(mirror_path)
    assert report.kind == "experiment"
    failed = [f"{r.sheet}!{r.ref}: {r.mismatches[:2]}" for r in report.results if not r.passed]
    assert not failed
    assert report.passed
    assert report.cells == 11631
    assert "Zhuravlyov Distances" in report.sheets()
    assert max(r.max_difference for r in report.results) < 1e-9


# ---------------------------------------------------------------- structure


def test_the_same_sheets_in_the_same_order(original: Any, mirror: Any) -> None:
    assert mirror.sheetnames == [RENAMED.get(name, name) for name in original.sheetnames]


def test_every_number_is_at_its_original_address(original: Any, mirror: Any) -> None:
    compared = 0
    for name, a, b in pairs(original, mirror):
        for x, y in cells(a, b):
            if in_run_block(name, x):
                continue
            where = f"{name}!{x.coordinate}"
            if is_number(x.value):
                assert is_number(y.value), f"{where}: {x.value!r} → {y.value!r}"
                assert abs(x.value - y.value) <= 1e-9, f"{where}: {x.value!r} → {y.value!r}"
                compared += 1
            else:
                assert not is_number(y.value), f"{where}: {x.value!r} → {y.value!r}"
    assert compared > 10_000


def test_the_text_differs_only_where_intended(original: Any, mirror: Any) -> None:
    found: dict[str, int] = {}
    for name, a, b in pairs(original, mirror):
        for x, y in cells(a, b):
            if in_run_block(name, x) or is_number(x.value) or is_number(y.value):
                continue
            if empty(x.value) and empty(y.value):
                continue
            if x.value != y.value:
                # text is never dropped: the mirror says something at every labelled cell
                assert not empty(y.value), f"{name}!{x.coordinate}: {x.value!r} was left out"
                found[name] = found.get(name, 0) + 1
    assert set(found) <= set(TEXT_DIFFERENCES)
    for name, count in found.items():
        assert count <= TEXT_DIFFERENCES[name], f"{name}: {count} text cells differ"


def test_the_run_block_is_the_only_addition(mirror: Any) -> None:
    sheet = mirror["Parameters"]
    assert str(sheet["A39"].value).startswith("This run")
    labels = [sheet.cell(row, 1).value for row in range(40, 48)]
    assert all(labels)
    assert "A39:C39" in [str(r) for r in sheet.merged_cells.ranges]


def test_the_same_merged_ranges_panes_and_tabs(original: Any, mirror: Any) -> None:
    for name, a, b in pairs(original, mirror):
        merged_a = {str(r) for r in a.merged_cells.ranges}
        merged_b = {str(r) for r in b.merged_cells.ranges}
        if name == "Parameters":
            merged_b.discard("A39:C39")
        assert merged_a == merged_b, name
        assert a.freeze_panes == b.freeze_panes, name
        assert a.sheet_properties.tabColor.rgb == b.sheet_properties.tabColor.rgb, name
        assert a.sheet_view.showGridLines == b.sheet_view.showGridLines, name


def conditional_formats(sheet: Any) -> list[tuple[Any, ...]]:
    found = []
    for cell_range, rules in sheet.conditional_formatting._cf_rules.items():
        for rule in rules:
            dxf = rule.dxf
            found.append(
                (
                    str(cell_range.sqref),
                    rule.type,
                    rule.operator,
                    tuple(rule.formula),
                    dxf.fill.bgColor.rgb if dxf and dxf.fill else None,
                    dxf.font.color.rgb if dxf and dxf.font and dxf.font.color else None,
                )
            )
    return found


def test_the_same_conditional_formats(original: Any, mirror: Any) -> None:
    total = 0
    for name, a, b in pairs(original, mirror):
        assert conditional_formats(a) == conditional_formats(b), name
        total += len(conditional_formats(a))
    assert total > 20


def widths(sheet: Any) -> dict[str, tuple[float, bool]]:
    found = {}
    for dimension in sheet.column_dimensions.values():
        if dimension.width:
            first = dimension.min or 1
            for column in range(first, (dimension.max or first) + 1):
                letter = openpyxl.utils.get_column_letter(column)
                found[letter] = (round(dimension.width, 2), bool(dimension.hidden))
    return found


def test_the_same_column_widths(original: Any, mirror: Any) -> None:
    for name, a, b in pairs(original, mirror):
        wa, wb = widths(a), widths(b)
        different = {column for column in set(wa) | set(wb) if wa.get(column) != wb.get(column)}
        assert different <= WIDTH_DIFFERENCES.get(name, set()), f"{name}: {sorted(different)}"


def colour(value: Any) -> str | None:
    if value is None or value.type != "rgb" or not isinstance(value.rgb, str):
        return None  # a theme colour: the default ink
    return str(value.rgb[2:])


def style(cell: Any) -> tuple[Any, ...]:
    font, border, alignment = cell.font, cell.border, cell.alignment
    return (
        colour(cell.fill.fgColor) if cell.fill and cell.fill.fill_type else None,
        font.sz,
        bool(font.b),
        bool(font.i),
        colour(font.color) if font.color else None,
        font.u,
        cell.number_format,
        alignment.horizontal,
        alignment.vertical,
        bool(alignment.wrap_text),
        float(alignment.indent or 0),
        "".join(
            (s.style or "-")[0] for s in (border.left, border.right, border.top, border.bottom)
        ),
    )


def test_the_same_cell_styles(original: Any, mirror: Any) -> None:
    compared = 0
    for name, a, b in pairs(original, mirror):
        # the style of an untouched cell (a detached cell: asking the sheet would create one)
        blank = (style(Cell(a, row=1, column=1)), style(Cell(b, row=1, column=1)))
        different = set()
        for x, y in cells(a, b):
            if in_run_block(name, x) or not (x.has_style or y.has_style):
                continue
            compared += 1
            sx, sy = style(x), style(y)
            if sx == sy or (sx, sy) == blank:
                continue
            # the font size of an empty cell without fill or border is not visible
            changed = {k for k in range(len(sx)) if sx[k] != sy[k]}
            if empty(x.value) and empty(y.value) and changed == {1}:
                continue
            different.add(x.coordinate)
        assert different <= STYLE_DIFFERENCES.get(name, set()), f"{name}: {sorted(different)}"
    assert compared > 15_000
