"""Template replication: the author's HAG and meta-algorithm template workbooks (ADR-001, ADR-029).

The workbooks are byte-for-byte copies in ``tests/data/templates`` (the template project is not
available in CI); on the author's machine a test checks that the copies still equal the originals.
"""

import hashlib
import json
import warnings
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pytest

from context_synthetic_recognition.config import CentreMode, HAGConfig, preset
from context_synthetic_recognition.core.hag import StopReason, hag
from context_synthetic_recognition.core.meta import K2_DECISION, meta_classify
from context_synthetic_recognition.notation import feature_name
from context_synthetic_recognition.services.validation import validate_workbook
from context_synthetic_recognition.services.validation.templates import (
    read_hag_template,
    read_meta_template,
)

from ..conftest import GOLDEN_VALUES, HAG_TEMPLATE, META_TEMPLATE, TEMPLATE_PROJECT

pytestmark = pytest.mark.golden


def _open(path: Path) -> Any:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return openpyxl.load_workbook(path, data_only=True)


@pytest.fixture(scope="module")
def hag_book() -> Iterator[Any]:
    book = _open(HAG_TEMPLATE)
    yield book
    book.close()


@pytest.fixture(scope="module")
def meta_book() -> Iterator[Any]:
    book = _open(META_TEMPLATE)
    yield book
    book.close()


def _set(tuplam: tuple[int, ...]) -> str:
    return "{" + ", ".join(feature_name(u) for u in tuplam) + "}"


@pytest.mark.parametrize(
    ("copy", "original"),
    [
        (HAG_TEMPLATE, Path("hag-algorithm") / HAG_TEMPLATE.name),
        (META_TEMPLATE, Path("meta-algorithm") / META_TEMPLATE.name),
    ],
)
def test_the_copies_equal_the_originals(copy: Path, original: Path) -> None:
    source = TEMPLATE_PROJECT / "resources" / "experiments" / original
    if not source.is_file():
        pytest.skip("the template project is not available (CI)")
    digest = [hashlib.sha256(path.read_bytes()).hexdigest() for path in (copy, source)]
    assert digest[0] == digest[1]


# ---------------------------------------------------------------- HAG template


def test_the_hag_template_is_reproduced(hag_book: Any) -> None:
    C, w, y, _ = read_hag_template(hag_book)
    result = hag(C, w, y)
    assert _set(result.tuplam) == "{x₃, x₆, x₁₃, x₄, x₉}"
    assert result.stop is StopReason.KAPPA
    latent = np.array(
        [
            [hag_book["Dataset for Meta-algorithm"].cell(r, c).value for c in range(10, 14)]
            for r in range(3, 13)
        ]
    )
    assert np.abs(result.latent - latent).max() <= 1e-12
    # θ/γ of every candidate of every iteration (column O of the candidate blocks)
    for j, it in enumerate(result.iterations, start=1):
        sheet = hag_book[f"Greedy upon Weight ({j}-Latent)"]
        cells = [sheet.cell(41 + 15 * c, 15).value for c in range(it.candidates.size)]
        assert np.abs(np.array(cells) - it.ratio).max() <= 1e-12
        assert sheet.cell(30 + 15 * it.candidates.size, 15).value == pytest.approx(
            it.crit, abs=1e-12
        )


def test_the_hag_template_validation_report() -> None:
    report = validate_workbook(HAG_TEMPLATE)
    assert report.passed, [r.mismatches[:3] for r in report.results if not r.passed]
    assert report.kind == "hag-template"
    assert report.cells == 6131
    assert max(r.max_difference for r in report.results) < 1e-14
    assert report.notes[0] == "SET = {x₃, x₆, x₁₃, x₄, x₉}; stop: |TUPLAM| = ϰ"


def test_the_switch_variants_on_the_template_data(hag_book: Any) -> None:
    with GOLDEN_VALUES.open(encoding="utf-8") as handle:
        golden = json.load(handle)["tpl"]
    C, w, y, _ = read_hag_template(hag_book)
    for variant in golden:
        config = HAGConfig(centres=CentreMode(variant["cent"]), step4_passes=variant["pas"])
        assert _set(hag(C, w, y, config).tuplam) == variant["T"], variant["label"]


def test_the_article_setting_does_not_reproduce_the_template() -> None:
    report = validate_workbook(HAG_TEMPLATE, preset("article"))
    assert not report.passed
    assert report.notes[0].startswith("SET = {x₃, x₁, x₅, x₉, x₆}")


# ---------------------------------------------------------------- meta-algorithm template


def test_the_meta_template_is_reproduced(meta_book: Any) -> None:
    description, query = read_meta_template(meta_book)
    assert query.tolist() == [4, 0, 7, 2, 0]
    result = meta_classify(description, query)
    assert result.decisions.tolist() == [K2_DECISION]
    steps = result.steps(0)
    assert steps.b1(4).tolist() == []  # B1(a₄) = ∅
    assert steps.b2(4).tolist() == [8, 9]  # B2(a₄) = {S₉, S₁₀}
    assert meta_book["Meta-algorithm"]["C223"].value == "Class 2"


def test_the_meta_template_validation_report() -> None:
    report = validate_workbook(META_TEMPLATE)
    assert report.passed
    assert report.kind == "meta-template"
    assert report.notes == ("B1(a₄) = ∅, B2(a₄) = {S₉, S₁₀}; class: Class 2",)


def test_the_hag_template_feeds_the_meta_template(hag_book: Any, meta_book: Any) -> None:
    # the latent features of the HAG template are the dᵢ of the meta-algorithm template
    C, w, y, _ = read_hag_template(hag_book)
    description, _ = read_meta_template(meta_book)
    result = hag(C, w, y)
    assert np.abs(result.latent - description.latent).max() <= 1e-12
    assert np.array_equal(result.class_index, description.class_index)


# ---------------------------------------------------------------- Template Deviations (experiment)


@pytest.fixture(scope="module")
def experiment_book() -> Iterator[Any]:
    from ..conftest import WORKBOOK

    book = _open(WORKBOOK)
    yield book
    book.close()


def test_the_template_deviations_tables(hag_book: Any, experiment_book: Any) -> None:
    """Rows 11–22, 32–41 and 47–50: the template data under both calculations."""
    sheet = experiment_book["Template Deviations"]
    C, w, y, labels = read_hag_template(hag_book)
    running = hag(C, w, y)
    final = hag(C, w, y, HAGConfig(centres=CentreMode.FINAL))
    first_running, first_final = running.iterations[0], final.iterations[0]
    assert np.array_equal(first_running.candidates, first_final.candidates)
    for row, u in enumerate(first_running.candidates, start=11):
        scan = final.scan(0, int(u))
        expected = [
            first_running.theta[row - 11],
            first_running.gamma[row - 11],
            first_running.ratio[row - 11],
            first_final.theta[row - 11],
            first_final.gamma[row - 11],
            first_final.ratio[row - 11],
            scan.centre1[0],
            scan.centre2[0],
        ]
        assert sheet.cell(row, 1).value == feature_name(int(u))
        cells = [sheet.cell(row, c).value for c in range(2, 10)]
        assert np.abs(np.array(cells) - np.array(expected)).max() <= 1e-12
    # r₁ on the template data (q = x₆): R + η_q, one pass (article), two passes (template)
    assert first_running.q == 5
    once = running.scan(0, 5)
    for t in range(10):
        cells = [sheet.cell(32 + t, c).value for c in range(2, 7)]
        expected = [once.b[t], once.majorized[t], first_running.latent[t], first_running.latent[t]]
        assert np.abs(np.array(cells[:4]) - np.array(expected)).max() <= 1e-12
        assert cells[4] == labels[t]
    # the SET of every switch setting
    for row, (centres, passes) in enumerate(
        [
            (CentreMode.RUNNING, 2),
            (CentreMode.RUNNING, 1),
            (CentreMode.FINAL, 2),
            (CentreMode.FINAL, 1),
        ],
        start=47,
    ):
        result = hag(C, w, y, HAGConfig(centres=centres, step4_passes=passes))
        assert sheet.cell(row, 4).value == _set(result.tuplam)
